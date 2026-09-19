"""
readers.py — turn an SI/BL attachment (any supported format) into plain text.

Supports: .txt, .xlsx, .docx, .pdf (text-layer PDFs).

Every reader either returns a non-empty string or raises ReaderError with a
short machine-readable `reason` so the pipeline can classify the failure
consistently:

    reason="unreadable"   file is corrupt, unsupported, or has no extractable
                          text at all (e.g. a scanned/image-only PDF with no
                          OCR pass applied)
"""
from pathlib import Path


class ReaderError(Exception):
    def __init__(self, reason, detail=""):
        self.reason = reason
        self.detail = detail
        super().__init__(f"{reason}: {detail}")


def read_txt(raw_bytes):
    try:
        text = raw_bytes.decode("utf-8", errors="replace")
    except Exception as exc:  # pragma: no cover - decode() with errors="replace" won't raise
        raise ReaderError("unreadable", str(exc))
    if not text.strip():
        raise ReaderError("unreadable", "empty text file")
    return text


def read_xlsx(raw_bytes):
    import io
    import openpyxl

    try:
        wb = openpyxl.load_workbook(io.BytesIO(raw_bytes), data_only=True)
    except Exception as exc:
        raise ReaderError("unreadable", f"could not open workbook: {exc}")

    lines = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                # Most rows in these sheets are "LABEL | VALUE" pairs.
                lines.append(" | ".join(cells))
    text = "\n".join(lines)
    if not text.strip():
        raise ReaderError("unreadable", "workbook has no readable cell content")
    return text


def read_docx(raw_bytes):
    import io
    from docx import Document

    try:
        doc = Document(io.BytesIO(raw_bytes))
    except Exception as exc:
        raise ReaderError("unreadable", f"could not open document: {exc}")

    lines = []
    for p in doc.paragraphs:
        if p.text.strip():
            lines.append(p.text.strip())
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(cells))
    text = "\n".join(lines)
    if not text.strip():
        raise ReaderError("unreadable", "document has no readable text or tables")
    return text


def read_pdf(raw_bytes):
    import io
    import pdfplumber

    try:
        with pdfplumber.open(io.BytesIO(raw_bytes)) as pdf:
            pages = [(page.extract_text() or "") for page in pdf.pages]
    except Exception as exc:
        # Covers corrupt/malformed PDFs (e.g. missing xref/root object).
        raise ReaderError("unreadable", f"corrupt or unparsable PDF: {exc}")

    text = "\n".join(pages)
    if not text.strip():
        # Almost certainly a scanned/image-only PDF. No OCR pass is applied
        # yet, so this is honestly reported as unreadable rather than guessed.
        raise ReaderError("unreadable", "no extractable text (likely a scanned/image PDF)")
    return text


_READERS = {
    ".txt": read_txt,
    ".xlsx": read_xlsx,
    ".docx": read_docx,
    ".pdf": read_pdf,
}


def read_attachment(inbox, att_path):
    """Read an attachment (given its path as it appears in email['attachments'])
    through the loader's `inbox`, dispatching on file extension.

    Returns plain text. Raises ReaderError with .reason in
    {"unreadable"} on any failure, or if the extension isn't supported.
    """
    ext = Path(att_path).suffix.lower()
    reader = _READERS.get(ext)
    if reader is None:
        raise ReaderError("unreadable", f"unsupported attachment type: {ext or '(none)'}")

    try:
        raw_bytes = inbox.read_bytes(att_path)
    except Exception as exc:
        raise ReaderError("unreadable", f"could not read file: {exc}")

    return reader(raw_bytes)
